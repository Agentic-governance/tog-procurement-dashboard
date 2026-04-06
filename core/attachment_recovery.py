#!/usr/bin/env python3
"""
core/attachment_recovery.py — Attachment Fetch, Parse & Recovery
Extracted from Night 9 / Night 14 / Day 1 Sprint

Handles:
- Attachment download (httpx)
- PDF/Word/Excel/HTML text extraction
- Failure taxonomy and tracking
- Priority-based batch processing
- Dead link analysis
"""

import sqlite3
import json
import hashlib
import re
from datetime import datetime

__all__ = [
    'create_failure_tables',
    'fetch_and_parse_attachment',
    'batch_recover_attachments',
    'classify_failure',
    'get_failure_taxonomy',
    'get_dead_link_analysis',
    'get_attachment_stats',
]

# Lazy imports for optional deps
_httpx = None
_pdfplumber = None
_docx = None
_openpyxl = None


def _ensure_deps():
    """Lazy-load optional dependencies."""
    global _httpx, _pdfplumber, _docx, _openpyxl
    if _httpx is None:
        try:
            import httpx
            _httpx = httpx
        except ImportError:
            pass
    if _pdfplumber is None:
        try:
            import pdfplumber
            _pdfplumber = pdfplumber
        except ImportError:
            pass
    if _docx is None:
        try:
            from docx import Document
            _docx = Document
        except ImportError:
            pass
    if _openpyxl is None:
        try:
            import openpyxl
            _openpyxl = openpyxl
        except ImportError:
            pass


def create_failure_tables(conn):
    """Create attachment_failures tracking table."""
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS attachment_failures (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        attachment_asset_id INTEGER NOT NULL,
        failure_type TEXT NOT NULL,
        failure_stage TEXT,
        retry_count INTEGER DEFAULT 0,
        recoverable_flag INTEGER DEFAULT 1,
        notes TEXT,
        created_at TEXT DEFAULT (datetime('now'))
    )''')
    conn.commit()


def classify_failure(error_message):
    """Classify an error into a failure type."""
    msg = str(error_message).lower()
    if '404' in msg:
        return 'http_404'
    elif '403' in msg:
        return 'http_403'
    elif '500' in msg or '502' in msg or '503' in msg:
        return 'server_error'
    elif 'timeout' in msg or 'timed out' in msg:
        return 'timeout'
    elif 'ssl' in msg or 'certificate' in msg:
        return 'ssl_error'
    elif 'connect' in msg:
        return 'connection_error'
    else:
        return 'unknown'


def fetch_and_parse_attachment(asset_id, url, asset_type, conn, procurement_item_id=None):
    """Download and extract text from a single attachment."""
    _ensure_deps()
    if _httpx is None:
        return {'status': 'error', 'error': 'httpx not available'}

    c = conn.cursor()
    if procurement_item_id is None:
        c.execute('SELECT procurement_item_id FROM attachment_assets WHERE id = ?', (asset_id,))
        row = c.fetchone()
        procurement_item_id = row[0] if row else None

    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}

    try:
        resp = _httpx.get(url, headers=headers, timeout=20, verify=False, follow_redirects=True)
        resp.raise_for_status()
        content = resp.content
        content_hash = hashlib.md5(content).hexdigest()

        text = ''
        confidence = 0.0

        if asset_type == 'pdf' and _pdfplumber:
            import io
            try:
                with _pdfplumber.open(io.BytesIO(content)) as pdf:
                    pages = [p.extract_text() for p in pdf.pages if p.extract_text()]
                    text = '\n'.join(pages)
                    confidence = 0.9 if text.strip() else 0.1
            except Exception:
                pass

        elif asset_type == 'word' and _docx:
            import io
            try:
                doc = _docx(io.BytesIO(content))
                text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
                confidence = 0.85 if text.strip() else 0.1
            except Exception:
                pass

        elif asset_type == 'excel' and _openpyxl:
            import io
            try:
                wb = _openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                rows_text = []
                for ws in wb.worksheets:
                    for r in ws.iter_rows(values_only=True):
                        rs = ' '.join(str(c2) for c2 in r if c2 is not None)
                        if rs.strip():
                            rows_text.append(rs)
                text = '\n'.join(rows_text)
                confidence = 0.7 if text.strip() else 0.1
            except Exception:
                pass

        elif asset_type == 'html_page':
            html = content.decode('utf-8', errors='ignore')
            text = re.sub(r'<[^>]+>', ' ', html)
            text = re.sub(r'\s+', ' ', text).strip()
            confidence = 0.6 if len(text) > 100 else 0.2

        else:
            return {'status': 'skipped', 'reason': f'unsupported: {asset_type}'}

        if text.strip():
            c.execute('SELECT id FROM attachment_parsed WHERE asset_id = ?', (asset_id,))
            if c.fetchone():
                c.execute('''UPDATE attachment_parsed SET text_content=?, extraction_method=?,
                             extraction_confidence=?, text_length=?, parsed_at=datetime('now')
                             WHERE asset_id=?''',
                          (text[:50000], f'{asset_type}_parse', confidence, len(text), asset_id))
            else:
                c.execute('''INSERT INTO attachment_parsed
                             (asset_id, procurement_item_id, text_content, extraction_method,
                              extraction_confidence, text_length, parsed_at)
                             VALUES (?,?,?,?,?,?,datetime('now'))''',
                          (asset_id, procurement_item_id, text[:50000],
                           f'{asset_type}_parse', confidence, len(text)))

            c.execute('''UPDATE attachment_assets SET parse_status='parsed',
                         content_hash=?, text_length=?, text_extracted_flag=1,
                         extraction_confidence=?, fetched_at=datetime('now')
                         WHERE id=?''',
                      (content_hash, len(text), confidence, asset_id))
            return {'status': 'parsed', 'text_length': len(text), 'confidence': confidence}
        else:
            c.execute('''UPDATE attachment_assets SET parse_status='empty',
                         content_hash=?, text_length=0, fetched_at=datetime('now')
                         WHERE id=?''', (content_hash, asset_id))
            return {'status': 'empty'}

    except Exception as e:
        err_msg = str(e)[:500]
        failure_type = classify_failure(err_msg)

        c.execute('''INSERT INTO attachment_failures
                     (attachment_asset_id, failure_type, failure_stage, notes)
                     VALUES (?, ?, 'fetch', ?)''',
                  (asset_id, failure_type, err_msg))

        c.execute('''UPDATE attachment_assets SET parse_status='error',
                     error_message=?, fetched_at=datetime('now')
                     WHERE id=?''', (err_msg, asset_id))
        return {'status': 'error', 'type': failure_type, 'error': err_msg[:200]}


def batch_recover_attachments(conn, max_items=250, min_score=30):
    """Batch process priority attachments."""
    c = conn.cursor()
    results = {'processed': 0, 'parsed': 0, 'errors': 0, 'skipped': 0}

    # Priority 1: bid_projects
    c.execute('''SELECT aa.id, aa.url, aa.asset_type, aa.procurement_item_id
                 FROM attachment_assets aa
                 JOIN bid_projects bp ON bp.item_id = aa.procurement_item_id
                 WHERE aa.parse_status IN ('pending')
                 AND aa.asset_type IN ('pdf', 'word', 'excel', 'html_page')
                 ORDER BY CASE WHEN bp.strategy_type='exploit' THEN 0 ELSE 1 END''')
    priority1 = c.fetchall()

    # Priority 2: high-score items
    c.execute('''SELECT aa.id, aa.url, aa.asset_type, aa.procurement_item_id
                 FROM attachment_assets aa
                 JOIN procurement_items pi ON pi.item_id = aa.procurement_item_id
                 WHERE aa.parse_status = 'pending'
                 AND aa.asset_type IN ('pdf', 'word', 'excel')
                 AND pi.important_v3_score >= ?
                 AND aa.procurement_item_id NOT IN (SELECT item_id FROM bid_projects)
                 ORDER BY pi.important_v3_score DESC
                 LIMIT ?''', (min_score, max_items))
    priority2 = c.fetchall()

    all_items = (priority1 + priority2)[:max_items]

    for idx, (asset_id, url, asset_type, item_id) in enumerate(all_items):
        if not url:
            results['skipped'] += 1
            continue

        result = fetch_and_parse_attachment(asset_id, url, asset_type, conn, item_id)
        results['processed'] += 1

        if result['status'] == 'parsed':
            results['parsed'] += 1
        elif result['status'] == 'error':
            results['errors'] += 1
        elif result['status'] == 'skipped':
            results['skipped'] += 1

        if idx % 10 == 0:
            conn.commit()

    conn.commit()
    return results


def get_failure_taxonomy(conn):
    """Get failure type distribution."""
    c = conn.cursor()
    c.execute('''SELECT failure_type, COUNT(*), COUNT(DISTINCT attachment_asset_id)
                 FROM attachment_failures
                 GROUP BY failure_type ORDER BY COUNT(*) DESC''')
    return [{'type': r[0], 'count': r[1], 'unique_assets': r[2]} for r in c.fetchall()]


def get_dead_link_analysis(conn):
    """Analyze dead links by domain."""
    c = conn.cursor()
    c.execute('''SELECT aa.url FROM attachment_failures af
                 JOIN attachment_assets aa ON aa.id = af.attachment_asset_id
                 WHERE af.failure_type = 'http_404' ''')
    urls = [r[0] for r in c.fetchall() if r[0]]

    from urllib.parse import urlparse
    domain_counts = {}
    for url in urls:
        try:
            domain = urlparse(url).netloc
            domain_counts[domain] = domain_counts.get(domain, 0) + 1
        except Exception:
            pass

    return sorted(domain_counts.items(), key=lambda x: -x[1])


def get_attachment_stats(conn):
    """Get comprehensive attachment statistics."""
    c = conn.cursor()
    c.execute('SELECT parse_status, COUNT(*) FROM attachment_assets GROUP BY parse_status')
    status_dist = dict(c.fetchall())

    c.execute('SELECT asset_type, COUNT(*) FROM attachment_assets GROUP BY asset_type ORDER BY COUNT(*) DESC')
    type_dist = dict(c.fetchall())

    c.execute('SELECT COUNT(*) FROM attachment_parsed')
    total_parsed = c.fetchone()[0]

    return {
        'status_distribution': status_dist,
        'type_distribution': type_dist,
        'total_parsed_text': total_parsed
    }
